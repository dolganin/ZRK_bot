import argparse
import os
import re
from datetime import UTC, date, datetime, time
from decimal import Decimal
from pathlib import Path
from urllib.parse import urlparse

import psycopg2
from openpyxl import Workbook


def normalize_database_url(raw_url: str) -> str:
    if raw_url.startswith("postgresql+asyncpg://"):
        return raw_url.replace("postgresql+asyncpg://", "postgresql://", 1)
    return raw_url


def make_output_path(explicit_path: str | None) -> Path:
    if explicit_path:
        return Path(explicit_path).expanduser().resolve()
    timestamp = datetime.now(UTC).strftime("%Y%m%d_%H%M%S")
    return Path.cwd() / f"zrk_export_{timestamp}.xlsx"


def sheet_name(name: str) -> str:
    safe = re.sub(r"[\[\]\*\?/\\:]", "_", name).strip()
    return safe[:31] or "sheet"


def excel_value(value):
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, datetime):
        if value.tzinfo is not None:
            return value.astimezone(UTC).replace(tzinfo=None)
        return value
    if isinstance(value, time):
        if value.tzinfo is not None:
            return value.replace(tzinfo=None)
        return value
    if isinstance(value, (datetime, date, time, int, float, bool)) or value is None:
        return value
    return str(value)


def fetch_table_names(cursor) -> list[str]:
    cursor.execute(
        """
        SELECT table_name
        FROM information_schema.tables
        WHERE table_schema = 'public'
          AND table_type = 'BASE TABLE'
        ORDER BY table_name
        """
    )
    return [row[0] for row in cursor.fetchall()]


def fetch_table_rows(cursor, table_name: str):
    cursor.execute(f'SELECT * FROM "{table_name}" ORDER BY 1')
    columns = [desc[0] for desc in cursor.description]
    rows = cursor.fetchall()
    return columns, rows


def build_meta_rows(db_url: str, tables: list[tuple[str, int]]) -> list[list[str | int]]:
    parsed = urlparse(db_url)
    return [
        ["exported_at_utc", datetime.now(UTC).isoformat(timespec="seconds")],
        ["database", (parsed.path or "/").lstrip("/")],
        ["host", parsed.hostname or ""],
        ["port", parsed.port or ""],
        ["tables_exported", len(tables)],
        ["total_rows", sum(row_count for _, row_count in tables)],
        [],
        ["table_name", "row_count"],
        *[[name, row_count] for name, row_count in tables],
    ]


def fetch_codes_usage_rows(cursor):
    cursor.execute(
        """
        SELECT
            c.id,
            c.code,
            e.name AS event_name,
            c.points,
            c.is_income,
            c.created_at,
            c.starts_at,
            c.expires_at,
            c.max_uses,
            COUNT(uc.user_id) AS usage_count
        FROM codes c
        LEFT JOIN events e ON e.id = c.event_id
        LEFT JOIN user_codes uc ON uc.code_id = c.id
        GROUP BY
            c.id,
            c.code,
            e.name,
            c.points,
            c.is_income,
            c.created_at,
            c.starts_at,
            c.expires_at,
            c.max_uses
        ORDER BY c.id
        """
    )
    columns = [desc[0] for desc in cursor.description]
    rows = cursor.fetchall()
    return columns, rows


def append_sheet(workbook: Workbook, title: str, columns, rows) -> None:
    sheet = workbook.create_sheet(title=sheet_name(title))
    sheet.append(columns)
    for row in rows:
        sheet.append([excel_value(value) for value in row])
    sheet.freeze_panes = "A2"


def export_database(database_url: str, output_path: Path) -> Path:
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "summary"

    with psycopg2.connect(database_url) as conn:
        with conn.cursor() as cursor:
            table_names = fetch_table_names(cursor)
            exported_tables: list[tuple[str, int]] = []

            for table_name in table_names:
                columns, rows = fetch_table_rows(cursor, table_name)
                exported_tables.append((table_name, len(rows)))
                append_sheet(workbook, table_name, columns, rows)

            codes_usage_columns, codes_usage_rows = fetch_codes_usage_rows(cursor)

    for row in build_meta_rows(database_url, exported_tables):
        summary_sheet.append(row)

    append_sheet(workbook, "codes_usage", codes_usage_columns, codes_usage_rows)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return output_path


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Export all public PostgreSQL tables into a single XLSX workbook."
    )
    parser.add_argument(
        "-o",
        "--output",
        help="Path to output XLSX file. Default: ./zrk_export_<timestamp>.xlsx",
    )
    args = parser.parse_args()

    database_url = normalize_database_url((os.getenv("DATABASE_URL") or "").strip())
    if not database_url:
        raise SystemExit("DATABASE_URL is not set")

    output_path = make_output_path(args.output)
    export_database(database_url, output_path)
    print(output_path)


if __name__ == "__main__":
    main()
